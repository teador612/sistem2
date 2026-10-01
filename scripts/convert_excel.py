"""Excel geçmiş verisini web uygulamasının ortak JSON formatına dönüştürür."""
from __future__ import annotations

import argparse
import json
import re
import unicodedata
from datetime import datetime
from pathlib import Path

import openpyxl


def parse_number(value):
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip().replace(".", "").replace(",", "."))
    except ValueError:
        return None


def clean_text(value):
    return re.sub(r"\s+", " ", str(value or "").replace("�", "").strip())


def normalize_team(value):
    text = clean_text(value).casefold()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "", text)


def date_text(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    text = clean_text(value)
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return text[:10]


def score_result(score):
    score = clean_text(score)
    if not re.match(r"^\d+\s*[-–]\s*\d+$", score):
        return None
    return score.replace("–", "-").replace(" ", "")


def convert(source: Path, destination: Path):
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    output = []

    for row in sheet.iter_rows(min_row=9, values_only=True):
        if not row[30] or not row[34] or not row[35]:
            continue
        date = date_text(row[30])
        home = clean_text(row[34])
        away = clean_text(row[35])
        half_score = score_result(row[28])
        full_score = score_result(row[29])
        status = "finished" if full_score else "not_started"
        markets = {
            "ms1": parse_number(row[36]), "ms0": parse_number(row[37]), "ms2": parse_number(row[38]),
            "kgYes": parse_number(row[55]), "kgNo": parse_number(row[56]),
            "iy15Under": parse_number(row[64]), "iy15Over": parse_number(row[65]),
            "under25": parse_number(row[51]), "over25": parse_number(row[52]),
        }
        results = {
            "msResult": clean_text(row[20]), "kgResult": clean_text(row[17]),
            "iy15Result": clean_text(row[1]), "over25Result": clean_text(row[7]),
        }
        output.append({
            "id": f"excel-{date}-{clean_text(row[33])}-{normalize_team(home)}-{normalize_team(away)}",
            "date": date, "time": clean_text(row[33]), "league": clean_text(row[32]),
            "homeTeam": home, "awayTeam": away,
            "normalizedHomeTeam": normalize_team(home), "normalizedAwayTeam": normalize_team(away),
            "openingOdds": markets, "halfTimeScore": half_score, "fullTimeScore": full_score,
            "results": results, "status": status, "lastUpdated": datetime.now().isoformat(timespec="seconds"),
        })

    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(output, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"{len(output)} maç yazıldı: {destination}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    args = parser.parse_args()
    convert(args.source, args.destination)
